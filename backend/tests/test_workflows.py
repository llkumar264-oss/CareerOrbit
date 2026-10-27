import base64, hashlib, io, json, time
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import httpx
import pytest
from docx import Document
from sqlalchemy import select,func
from app import mail,providers
from app.db import now,SessionLocal
from app.models import *
from app.schemas import Preferences
from app.security import cipher
from app.matching import match,company_key
from app.outreach import write_email,scope_hash,validate_draft,resume_bytes
from app.worker import tick

def campaign_payload(**overrides):
    p={'query':'Python','types':['internship'],'locations':['Bangalore'],'send_start_hour':0,'send_end_hour':24,'timezone':'UTC','mode':'review','sources':{'greenhouse':['testboard']}}
    p.update(overrides)
    return {'name':'Data internships','preferences':p}
def fixture_row(i=0):
    return dict(source='greenhouse',external_id=str(i),company=f'Synthetic Company {i}',title='Python internship',description='Build Python data pipelines to help the analytics team process research datasets and improve reporting.',url=f'https://example.com/synthetic/{i}',location='Bangalore',classification='A',data={'type':'internship','application_method':'email','synthetic':True})
def populate(db,n=1,**pref_overrides):
    user=User(email='test@example.com',password_hash='unused');db.add(user);db.flush()
    profile=Profile(user_id=user.id,data={'full_name':'Alex Applicant','skills':['Python'],'interests':['data']},confirmed=True);db.add(profile)
    prefs=Preferences(**campaign_payload(**pref_overrides)['preferences']).model_dump(mode='json')
    c=Campaign(user_id=user.id,name='Synthetic test campaign',preferences=prefs,status='active');db.add(c);db.flush()
    content=b'%PDF-1.4 test attachment';path=Path(cipher_test_dir())/f'{uid()}.enc';path.write_bytes(cipher().encrypt(content))
    r=Resume(user_id=user.id,name='Data',filename='data.pdf',path=str(path),mime='application/pdf',sha256=hashlib.sha256(content).hexdigest(),text='Alex Applicant has Python data project experience.',approved=True);db.add(r)
    box=Mailbox(user_id=user.id,provider='gmail',address=user.email,tokens=cipher().encrypt(json.dumps({'access_token':'test','expires_at':time.time()+3600}).encode()).decode());db.add(box);db.flush()
    ds=[]
    for i in range(n):
        row=fixture_row(i);o=Opportunity(campaign_id=c.id,company_key=company_key(row['company']),score=90,eligible=True,explanation={},**row);ct=Contact(user_id=user.id,company_key=company_key(row['company']),email=f'hr{i}@example.com',source_url=row['url'],verification='valid');db.add_all([o,ct]);db.flush()
        subject,body=write_email(profile.data,o,r)
        d=Draft(campaign_id=c.id,opportunity_id=o.id,contact_id=ct.id,resume_id=r.id,subject=subject,body=body,status='approved',validation=[]);db.add(d);db.flush();ds.append(d)
    a=Approval(campaign_id=c.id,revision=c.revision,scope_hash=scope_hash(c,ds),draft_ids=[d.id for d in ds],max_messages=n);db.add(a);db.flush()
    for d in ds: db.add(QueueJob(draft_id=d.id,approval_id=a.id,mailbox_id=box.id))
    db.commit();return user,c,r,box,ds

def cipher_test_dir():
    from app.config import settings
    return settings.storage_dir

def test_auth_ownership_csrf_persistence(signed):
    r=signed.put('/api/profile',json={'full_name':'Alex','skills':['Python'],'confirmed':True});assert r.status_code==200
    c=signed.post('/api/campaigns',json=campaign_payload()).json()
    assert signed.get('/api/campaigns').json()[0]['id']==c['id']
    with SessionLocal() as db: assert db.scalar(select(Campaign)).name=='Data internships'
    assert signed.post('/api/campaigns/'+c['id']+'/pause',headers={'Origin':'https://evil.example'}).status_code==403
    signed.post('/api/auth/logout')
    signed.post('/api/auth/register',json={'email':'other@example.com','password':'A-secure-test-password-002'})
    assert signed.get('/api/campaigns/'+c['id']+'/opportunities').status_code==404
    assert signed.get('/api/campaigns').json()==[]

def test_password_and_logout(signed):
    assert signed.get('/api/auth/me').status_code==200
    signed.post('/api/auth/logout');assert signed.get('/api/auth/me').status_code==401
    assert signed.post('/api/auth/login',json={'email':'applicant@example.com','password':'bad-password-123'}).status_code==401
    assert signed.post('/api/auth/login',json={'email':'applicant@example.com','password':'A-secure-test-password-001'}).status_code==200

def test_resume_upload_parse_approve_download(signed):
    doc=Document();doc.add_paragraph('Alex Applicant — Python and SQL. University student working on data science projects.');buf=io.BytesIO();doc.save(buf)
    r=signed.post('/api/resumes',files={'file':('resume.docx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'name':'Data science'})
    assert r.status_code==200,r.text
    d=r.json();assert 'Python' in d['text'];assert not d['approved'];assert 'path' not in d
    assert signed.post('/api/resumes/'+d['id']+'/approve').json()['approved']
    assert signed.get('/api/resumes/'+d['id']+'/download').content==buf.getvalue()
    with SessionLocal() as db: assert not Path(db.get(Resume,d['id']).path).read_bytes().startswith(b'PK')
    assert signed.post('/api/resumes',files={'file':('bad.pdf',b'not a PDF')},data={'name':'bad'}).status_code==422
    assert signed.post('/api/resumes',files={'file':('x.exe',b'MZ')},data={'name':'bad'}).status_code==422

def test_pdf_extraction(signed):
    # A valid minimal PDF with a text stream, constructed without external fixtures.
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
    w=PdfWriter();page=w.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 50 700 Td (Alex Applicant Python data engineering project experience) Tj ET');page[NameObject('/Contents')]=w._add_object(stream)
    buf=io.BytesIO();w.write(buf)
    r=signed.post('/api/resumes',files={'file':('resume.pdf',buf.getvalue())},data={'name':'PDF'});assert r.status_code==200,r.text;assert 'Python' in r.json()['text']

def test_matching_100_synthetic_records():
    p=Preferences(**campaign_payload()['preferences']).model_dump(mode='json')
    scored=[match(fixture_row(i),p,{'skills':['Python'],'interests':['data']}) for i in range(100)]
    assert len(scored)==100 and all(0<=s<=100 and e for s,_,e in scored)
    op=fixture_row();op['location']='Berlin';assert not match(op,p,{'skills':['Python']})[2]
    op=fixture_row();op['data']['eligibility']={'sponsorship':False};assert not match(op,p,{'requires_sponsorship':True})[2]
    op=fixture_row();op['data']['compensation']={'min':1000,'currency':'INR','frequency':'month'};p['minimum']=25000;assert not match(op,p,{})[2]
    op['data']['compensation']['frequency']='year';assert match(op,p,{})[1]['components']['compensation']==.5
    p['allow_undisclosed']=False;assert not match(fixture_row(),p,{})[2]
    op=fixture_row();op['classification']='D';assert not match(op,p,{})[2]
    op=fixture_row();op['data']['deadline']='2000-01-01';assert not match(op,p,{})[2]

def test_research_dedup_and_outage(signed):
    c=signed.post('/api/campaigns',json=campaign_payload()).json();path='/api/campaigns/'+c['id']
    with patch('app.providers.greenhouse',return_value=[fixture_row(i) for i in range(100)]):
        assert signed.post(path+'/research').json()['added']==100
        assert signed.post(path+'/research').json()['added']==0
    assert len(signed.get(path+'/opportunities').json())==100
    with patch('app.providers.greenhouse',side_effect=providers.ProviderError('Provider returned HTTP 503')):
        r=signed.post(path+'/research');assert r.status_code==200 and len(r.json()['errors'])==1
    assert signed.get('/api/dashboard').json()['researched']==100

def test_adzuna_missing_credentials():
    with patch.object(providers.settings,'adzuna_app_key',''):
        with pytest.raises(providers.ProviderError,match='not configured'):providers.adzuna('Python','in','Bangalore')

def test_provider_adapter_contracts():
    with patch('app.providers.fetch_json',return_value={'jobs':[{'id':1,'title':'Python internship','absolute_url':'https://example.com/jobs/1','content':'<p>Build Python</p>','location':{'name':'Bangalore'}}]}) as f:
        rows=providers.greenhouse('board');assert rows[0]['classification']=='A';assert rows[0]['description']==' Build Python ';assert rows[0]['data']['application_method']=='portal'
    with pytest.raises(providers.ProviderError): providers.greenhouse('https://internal.example')

def test_personalization_quality_attachment(db):
    u,c,r,b,ds=populate(db);d=ds[0];o=db.get(Opportunity,d.opportunity_id);ct=db.get(Contact,d.contact_id)
    assert o.company in d.body and o.title in d.body and 'Python' in d.body
    assert not validate_draft(d,o,ct,r,{'confirmed':True})
    d.body=d.body.replace(o.company,'Wrong company');assert 'Company or position missing' in validate_draft(d,o,ct,r,{'confirmed':True})
    d.body='{{placeholder}}';assert any('Placeholder' in x for x in validate_draft(d,o,ct,r,{'confirmed':True}))
    ct.company_key='other';assert any('another company' in x for x in validate_draft(d,o,ct,r,{'confirmed':True}))
    ct.verification='invalid';assert any('not verified' in x for x in validate_draft(d,o,ct,r,{'confirmed':True}))
    msg=mail.build_mime(b.address,ct.email,'Application','Body',r.filename,r.mime,resume_bytes(r),'job')
    assert len(list(msg.iter_attachments()))==1
    assert list(msg.iter_attachments())[0].get_payload(decode=True)==resume_bytes(r)
    Path(r.path).write_bytes(b'corrupt');assert any('integrity' in x for x in validate_draft(d,o,ct,r,{'confirmed':True}))

def test_queue_send_persist_and_no_duplicate(db):
    u,c,r,b,ds=populate(db)
    with patch('app.mail.send',return_value=('provider-123','thread-123')) as send:
        assert tick(db);assert not tick(db);assert send.call_count==1
    m=db.scalar(select(Message));assert m.status=='sent' and m.provider_id=='provider-123'
    assert db.scalar(select(QueueJob)).status=='sent'
    with SessionLocal() as fresh: assert fresh.scalar(select(Message)).recipient=='hr0@example.com'

def test_pause_resume_and_changed_scope(db):
    u,c,r,b,ds=populate(db);c.status='paused';db.commit()
    with patch('app.mail.send') as send:tick(db);assert not send.called
    job=db.scalar(select(QueueJob));job.next_at=now();c.status='active';c.revision+=1;db.commit()
    with patch('app.mail.send') as send:tick(db);assert not send.called
    assert job.status=='failed' and 'scope' in job.error

def test_daily_quota(db):
    u,c,r,b,ds=populate(db,n=2,daily_limit=1)
    with patch('app.mail.send',return_value=('id','thread')) as send:
        tick(db);tick(db);assert send.call_count==1
    assert db.scalar(select(func.count()).select_from(Message))==1
    assert len(list(db.scalars(select(QueueJob).where(QueueJob.status=='pending'))))==1

def test_transient_retry_and_ambiguous_no_retry(db):
    populate(db)
    with patch('app.mail.send',side_effect=mail.SendError('rate limit',transient=True,retry_after=90)):
        tick(db)
    j=db.scalar(select(QueueJob));assert j.status=='pending' and j.next_at>now()
    j.next_at=now();db.commit()
    with patch('app.mail.send',side_effect=mail.SendError('timeout',uncertain=True)) as send:
        tick(db);tick(db);assert send.call_count==1
    assert j.status=='unknown' and db.scalar(select(func.count()).select_from(Message))==0

def test_suppression_blocks_dispatch(db):
    u,c,r,b,ds=populate(db);db.add(Suppression(user_id=u.id,email='hr0@example.com',reason='unsubscribe'));db.commit()
    with patch('app.mail.send') as send: tick(db);assert not send.called
    assert db.scalar(select(QueueJob)).error=='Recipient suppressed'

def test_expired_evidence_blocks_dispatch(db):
    populate(db);o=db.scalar(select(Opportunity));o.verified_at=now()-timedelta(days=2);db.commit()
    with patch('app.mail.send') as send: tick(db);assert not send.called
    assert 'older than' in db.scalar(select(QueueJob)).error

def test_oauth_refresh_tokens_encrypted(db):
    u,c,r,b,ds=populate(db)
    b.tokens=cipher().encrypt(json.dumps({'access_token':'old-secret','refresh_token':'refresh-secret','expires_at':0}).encode()).decode();db.commit()
    with patch('app.mail.token_request',return_value={'access_token':'new-secret','expires_at':time.time()+3600}):
        assert mail.access_token(db,b)=='new-secret'
    assert 'new-secret' not in b.tokens
    assert json.loads(cipher().decrypt(b.tokens.encode()))['refresh_token']=='refresh-secret'

def test_oauth_connection_state_and_missing_config(signed):
    with patch.object(mail.settings,'gmail_client_id',''):
        assert signed.post('/api/oauth/gmail/connect').status_code==502
    assert signed.get('/api/oauth/gmail/callback?state=invalid&code=test').status_code==400
    with patch.object(mail.settings,'gmail_client_id','test'),patch.object(mail.settings,'gmail_client_secret','test'):
        r=signed.post('/api/oauth/gmail/connect?read=true');assert r.status_code==200
        assert 'code_challenge_method=S256' in r.json()['url'] and 'gmail.readonly' in r.json()['url']

def test_reply_classification_followup_suppression(signed,db):
    user=db.scalar(select(User));profile=db.get(Profile,user.id);profile.confirmed=True;profile.data={'full_name':'Alex','skills':['Python']};db.commit()
    # Build test data for the signed-in user by reassigning the fixture owner safely.
    other,c,r,b,ds=populate(db);c.user_id=user.id;r.user_id=user.id;b.user_id=user.id
    ct=db.get(Contact,ds[0].contact_id);ct.user_id=user.id;db.commit()
    m=Message(user_id=user.id,campaign_id=c.id,draft_id=ds[0].id,company_key=ct.company_key,recipient=ct.email,provider_id='mail1',thread_id='thread1',sent_at=now()-timedelta(days=9));db.add(m);b.read_enabled=True;db.commit()
    assert len(signed.get('/api/follow-ups').json())==1
    with patch('app.mail.list_replies',return_value=[{'id':'reply1','thread_id':'thread1','from':ct.email,'body':'Please stop contacting me','in_reply_to':''}]):
        assert signed.post('/api/mailboxes/'+b.id+'/sync').json()['persisted']==1
        assert signed.post('/api/mailboxes/'+b.id+'/sync').json()['persisted']==0
    assert signed.get('/api/follow-ups').json()==[]
    assert db.scalar(select(Suppression)).reason=='unsubscribe'
    assert signed.get('/api/dashboard').json()['replies']==1
    assert mail.classify_reply('We invite you to an interview')=='interview'

def test_portal_not_email_application(signed,db):
    c=signed.post('/api/campaigns',json=campaign_payload()).json()
    row=fixture_row();row['data']['application_method']='portal'
    with patch('app.providers.greenhouse',return_value=[row]):signed.post('/api/campaigns/'+c['id']+'/research')
    o=signed.get('/api/campaigns/'+c['id']+'/opportunities').json()[0]
    ct=signed.post('/api/contacts',json={'opportunity_id':o['id'],'email':'hr@example.com','source_url':o['url']}).json()
    assert signed.post('/api/drafts',json={'opportunity_id':o['id'],'contact_id':ct['id']}).status_code==409
    assert signed.get('/api/dashboard').json()['portal_submitted']==0
    signed.put('/api/applications/'+o['id'],json={'status':'portal_submitted','notes':'Submitted by me through official portal'})
    assert signed.get('/api/dashboard').json()['portal_submitted']==1

def test_export_and_delete(signed):
    c=signed.post('/api/campaigns',json=campaign_payload()).json()
    assert signed.get('/api/account/export').json()['campaigns'][0]['id']==c['id']
    assert signed.delete('/api/account').status_code==200
    assert signed.get('/api/auth/me').status_code==401
    with SessionLocal() as db: assert db.scalar(select(func.count()).select_from(User))==0

def test_one_approved_followup_and_reply_cancels(db):
    from app.worker import follow_up_tick
    u,c,r,b,ds=populate(db)
    with patch('app.mail.send',return_value=('initial','thread')): tick(db)
    m=db.scalar(select(Message));m.sent_at=now()-timedelta(days=9)
    f=FollowUp(user_id=u.id,message_id=m.id,mailbox_id=b.id,campaign_revision=c.revision,subject='Re: application',body='Hello, following up on my original application. Please let me know if the opportunity is still available.',next_at=now());db.add(f);db.commit()
    with patch('app.mail.send',return_value=('follow-up','thread')) as send:
        assert follow_up_tick(db);assert not follow_up_tick(db);assert send.call_count==1
    assert f.status=='sent'
    assert db.scalar(select(func.count()).select_from(FollowUp))==1

def test_followup_stops_on_reply(db):
    from app.worker import follow_up_tick
    u,c,r,b,ds=populate(db)
    with patch('app.mail.send',return_value=('initial','thread')): tick(db)
    m=db.scalar(select(Message));f=FollowUp(user_id=u.id,message_id=m.id,mailbox_id=b.id,campaign_revision=c.revision,subject='Re: application',body='Follow-up',next_at=now());db.add(f)
    db.add(Reply(user_id=u.id,mailbox_id=b.id,message_id=m.id,provider_id='reply',category='unclassified',body='Please get back to us next month'));db.commit()
    with patch('app.mail.send') as send:follow_up_tick(db);assert not send.called
    assert f.status=='cancelled'

def test_send_provider_failures_and_payload(db):
    u,c,r,b,ds=populate(db);d=ds[0];ct=db.get(Contact,d.contact_id)
    class MockClient:
        def __init__(self,status): self.status=status
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def post(self,url,**kw):
            assert 'raw' in kw['json']
            from email import message_from_bytes
            msg=message_from_bytes(base64.urlsafe_b64decode(kw['json']['raw']))
            assert msg['To']==ct.email
            assert len([p for p in msg.walk() if p.get_content_disposition()=='attachment'])==1
            return httpx.Response(self.status,headers={'Retry-After':'120'},json={'id':'provider-id','threadId':'thread'})
    for status,transient,uncertain in [(401,False,False),(429,True,False),(503,False,True)]:
        with patch('app.mail.client',return_value=MockClient(status)):
            with pytest.raises(mail.SendError) as e:mail.send(db,b,d,ct,r,resume_bytes(r),'jobid')
        assert e.value.transient==transient and e.value.uncertain==uncertain
    with patch('app.mail.client',return_value=MockClient(200)):
        assert mail.send(db,b,d,ct,r,resume_bytes(r),'jobid')==('provider-id','thread')

def test_llm_cannot_invent_output():
    from app.llm import select_evidence
    from app.config import settings
    class LLMClient:
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def post(self,*a,**kw): return httpx.Response(200,json={'choices':[{'message':{'content':'{"sentence_index":999,"skill_indexes":[0]}'}}]})
    with patch.object(settings,'llm_base_url','https://llm.example/v1'),patch.object(settings,'llm_api_key','test'),patch.object(settings,'llm_model','test'),patch('app.llm.client',return_value=LLMClient()):
        with pytest.raises(providers.ProviderError,match='invalid structured'):select_evidence(['Source text'],['Python'],'Intern')

def test_approval_exact_scope_and_edits_revoke(signed,db):
    user=db.scalar(select(User));other,c,r,b,ds=populate(db)
    c.user_id=user.id;r.user_id=user.id;b.user_id=user.id;ct=db.get(Contact,ds[0].contact_id);ct.user_id=user.id
    p=db.get(Profile,user.id);p.confirmed=True;p.data={'full_name':'Alex Applicant','skills':['Python']};db.commit()
    payload={'draft_ids':[ds[0].id],'mailbox_id':b.id,'confirmation':'I authorize these recipients, messages, attachments and campaign limits'}
    approved=signed.post('/api/campaigns/'+c.id+'/approve',json=payload);assert approved.status_code==200,approved.text
    assert signed.put('/api/drafts/'+ds[0].id,json={'subject':ds[0].subject,'body':ds[0].body+'\nThank you.'}).status_code==200
    db.expire_all();assert db.get(Campaign,c.id).status=='paused'
    db.get(Campaign,c.id).status='active';db.commit()
    with patch('app.mail.send') as send:tick(db);assert not send.called
    assert db.scalar(select(QueueJob)).status=='failed'

def test_complete_mocked_application_workflow(signed,db):
    from urllib.parse import urlparse,parse_qs
    # Only provider boundaries are mocked. API, database, parsing and dispatch run normally.
    assert signed.put('/api/profile',json={'full_name':'Alex Applicant','skills':['Python'],'interests':['data'],'confirmed':True}).status_code==200
    doc=Document();doc.add_paragraph('Alex Applicant: Python data pipeline projects and university research experience.');buf=io.BytesIO();doc.save(buf)
    r=signed.post('/api/resumes',files={'file':('resume.docx',buf.getvalue())},data={'name':'Data science'}).json()
    signed.post('/api/resumes/'+r['id']+'/approve')
    with patch.object(mail.settings,'gmail_client_id','test-client'),patch.object(mail.settings,'gmail_client_secret','test-secret'):
        auth=signed.post('/api/oauth/gmail/connect?read=true').json()
        state=parse_qs(urlparse(auth['url']).query)['state'][0]
        with patch('app.mail.token_request',return_value={'access_token':'test-token','refresh_token':'test-refresh','expires_at':time.time()+3600}),patch('app.mail.identity',return_value='applicant@example.com'):
            callback=signed.get('/api/oauth/gmail/callback',params={'state':state,'code':'mock-authorization-code'},follow_redirects=False)
            assert callback.status_code==307,callback.text
        # One-use state cannot be replayed.
        assert signed.get('/api/oauth/gmail/callback',params={'state':state,'code':'mock-code'}).status_code==400
    box=signed.get('/api/mailboxes').json()[0];assert 'tokens' not in box
    c=signed.post('/api/campaigns',json=campaign_payload()).json()
    with patch('app.providers.greenhouse',return_value=[fixture_row()]):assert signed.post('/api/campaigns/'+c['id']+'/research').json()['added']==1
    op=signed.get('/api/campaigns/'+c['id']+'/opportunities').json()[0]
    ct=signed.post('/api/contacts',json={'opportunity_id':op['id'],'email':'hr@example.com','source_url':op['url']}).json()
    with patch('app.providers.hunter_verify',return_value={'status':'valid','score':100,'result':'deliverable'}):
        assert signed.post('/api/contacts/'+ct['id']+'/verify').json()['verification']=='valid'
    dr=signed.post('/api/drafts',json={'opportunity_id':op['id'],'contact_id':ct['id']});assert dr.status_code==200,dr.text
    draft=dr.json();assert draft['resume_id']==r['id'];assert draft['validation']==[]
    approved=signed.post('/api/campaigns/'+c['id']+'/approve',json={'draft_ids':[draft['id']],'mailbox_id':box['id'],'confirmation':'I authorize these recipients, messages, attachments and campaign limits'})
    assert approved.status_code==200,approved.text
    with patch('app.mail.send',return_value=('provider-accepted','real-thread-shaped-id')) as transport:assert tick(db);assert transport.call_count==1
    assert signed.get('/api/dashboard').json()['sent']==1
    assert signed.get('/api/dashboard').json()['portal_submitted']==0
    with patch('app.mail.list_replies',return_value=[{'id':'provider-reply','thread_id':'real-thread-shaped-id','from':'hr@example.com','body':'We invite you to an interview','in_reply_to':''}]):
        assert signed.post('/api/mailboxes/'+box['id']+'/sync').json()['persisted']==1
    dashboard=signed.get('/api/dashboard').json();assert dashboard['replies']==1;assert dashboard['positive']==1
    assert signed.get('/api/replies').json()[0]['category']=='interview'
    # A classified email does not invent a scheduled interview date.
    assert dashboard['interviews']==0
    assert signed.delete('/api/account').status_code==200

def test_natural_language_campaign_is_editable_not_authorization(signed):
    text='I am a Python and data science student. Find 100 companies in Bangalore, Hyderabad, and Singapore offering paid internships, preferably above ₹25,000 per month. Contact 15 unique approved companies per day.'
    r=signed.post('/api/campaigns/parse',json={'text':text});assert r.status_code==200,r.text
    p=r.json()['preferences'];assert p['query']=='python data science';assert p['minimum']==25000;assert p['frequency']=='month';assert p['daily_limit']==15;assert p['target_companies']==100;assert p['mode']=='draft_only'
    assert p['locations']==['Bangalore','Hyderabad','Singapore']
    assert signed.get('/api/campaigns').json()==[]

def test_currency_comparison_preserves_frequency_and_provenance():
    from app.compensation import compare_compensation,daily_rate
    p=Preferences(query='Python',currency='INR',frequency='month',minimum=25000).model_dump(mode='json')
    comp={'min':1000,'currency':'SGD','frequency':'month'}
    daily_rate.cache_clear()
    with patch('app.compensation.fetch_json',return_value={'rates':{'INR':60},'date':'2026-10-27'}):
        result=compare_compensation(comp,p)
    assert result['min']==60000 and result['fx']['date']=='2026-10-27'
    assert comp=={'min':1000,'currency':'SGD','frequency':'month'}
    assert compare_compensation({**comp,'frequency':'year'},p) is None
    daily_rate.cache_clear()
    with patch('app.compensation.fetch_json',side_effect=providers.ProviderError('outage')):assert compare_compensation(comp,p) is None
